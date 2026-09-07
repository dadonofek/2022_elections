import openpyxl, json, re, config
wb = openpyxl.load_workbook(config.MATCHING_XLSX, data_only=False)

ws = wb['קלפיות']
rows = list(ws.iter_rows(values_only=True))
hdr = rows[1]
idx = {h: i for i, h in enumerate(hdr)}
party_cols = [h for h in hdr if isinstance(h, str) and h.startswith('מפלגה ')]

stations = []
for r in rows[2:]:
    if not r or r[0] is None: continue
    d = {
        'kalpi': str(r[idx['מספר קלפי']]),
        'barzel': r[idx['מספר ברזל']],
        'site': r[idx['שם אתר']],
        'address': r[idx['כתובת מועמדת']],
        'confidence': r[idx['רמת ביטחון']],
        'eligible': r[idx['בעלי זכות']],
        'voters': r[idx['מצביעים']],
        'invalid': r[idx['פסולים']],
        'valid': r[idx['כשרים']],
        'parties': {h.replace('מפלגה ', ''): (r[idx[h]] or 0) for h in party_cols},
    }
    stations.append(d)

wsa = wb['אתרים']
arows = list(wsa.iter_rows(values_only=True))
sites = []
for r in arows[2:]:
    if not r or r[0] is None: continue
    sites.append({'id': r[0], 'name': r[1], 'address': r[2], 'n_kalpi': r[3],
                  'kalpi_list': r[4], 'confidence': r[6], 'method': r[7], 'source': r[8]})

json.dump({'stations': stations, 'sites': sites}, open('data/raw.json','w'), ensure_ascii=False, indent=1)
print('stations', len(stations), 'sites', len(sites))
print('party cols', len(party_cols), party_cols[:5])
addrs = sorted({s['address'] for s in stations if s['address']})
print('unique addresses', len(addrs))
print('missing address', sum(1 for s in stations if not s['address']))
tot_elig = sum(s['eligible'] or 0 for s in stations); tot_vot = sum(s['voters'] or 0 for s in stations)
print('eligible', tot_elig, 'voters', tot_vot, 'turnout', round(100*tot_vot/tot_elig,2))
